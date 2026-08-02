"""Create the editable DOCX report; the document skill renderer emits the PDF."""

from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "report"
OUTPUTS = ROOT / "outputs"
FIGURES = OUTPUTS / "figures"
MODELS = ROOT / "models"
REPORT.mkdir(exist_ok=True)

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(23, 51, 70)
MUTED = RGBColor(90, 103, 112)
GOLD = RGBColor(183, 101, 61)


def set_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tc_pr = tc.get_or_add_tcPr(); margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar"); tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}"); margins.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); run._r.addnext(fld)


def configure(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8); normal.paragraph_format.line_spacing = 1.333
    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10), ("Heading 2", 13, BLUE, 12, 6), ("Heading 3", 12, DARK_BLUE, 8, 4)):
        style = doc.styles[style_name]; style.font.name = "Calibri"; style.font.size = Pt(size)
        style.font.color.rgb = color; style.font.bold = True
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = "CIS6005 Computational Intelligence | MEDA Virtual Sensor Recovery"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs: set_font(run, size=8.5, color=MUTED)
    add_page_number(section.footer.paragraphs[0])
    for run in section.footer.paragraphs[0].runs: set_font(run, size=8.5, color=MUTED)


def para(doc, text, bold_lead=None):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):]); set_font(r)
    else:
        r = p.add_run(text); set_font(r)
    return p


def add_figure(doc, path, caption, width=6.25):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption); set_font(r, size=9, color=MUTED, italic=True)
    c.paragraph_format.space_after = Pt(8); c.paragraph_format.keep_with_next = False


def add_metric_table(doc, comparison):
    selected = comparison[comparison.split.eq("validation")].sort_values("mse")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"; set_table_widths(table, [1.5, 1.1, 1.1, 1.1, 1.7])
    headers = ["Model", "MSE", "RMSE", "MAE", "R2"]
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label; shade(cell, "F2F4F7")
        for run in cell.paragraphs[0].runs: set_font(run, size=9, bold=True)
    for _, row in selected.iterrows():
        cells = table.add_row().cells
        values = [row.model, f"{row.mse:.3f}", f"{row.rmse:.3f}", f"{row.mae:.3f}", f"{row.r2:.3f}"]
        for cell, value in zip(cells, values):
            cell.text = str(value)
            for run in cell.paragraphs[0].runs: set_font(run, size=8.5)


def report_word_count(doc):
    count, active = 0, True
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "References": active = False
        if active: count += len(re.findall(r"\b[\w'-]+\b", text))
    return count


def main():
    metadata = json.loads((MODELS / "metadata_v2.json").read_text(encoding="utf-8"))
    comparison = pd.read_csv(OUTPUTS / "model_comparison_v2.csv")
    raw = pd.read_parquet(ROOT / "data" / "processed" / "stratified_training_sample.parquet")
    missing = raw.isna().mean().mul(100).sort_values(ascending=False)
    desc = raw.PRESSURE.describe()
    final = metadata["final_test_metrics"]
    validation = comparison[comparison.split.eq("validation")].sort_values("mse").iloc[0]

    doc = Document(); configure(doc)
    # Editorial cover override for the narrative_proposal preset.
    for _ in range(4): doc.add_paragraph()
    kicker = doc.add_paragraph(); kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(kicker.add_run("COMPUTATIONAL INTELLIGENCE PROJECT"), size=10.5, color=GOLD, bold=True)
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(title.add_run("MEDA Atmospheric Pressure\nVirtual Sensor Recovery"), size=28, color=NAVY, bold=True)
    subtitle = doc.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(subtitle.add_run("Temporally honest machine learning for NASA Perseverance rover observations"), size=14, color=DARK_BLUE)
    for _ in range(3): doc.add_paragraph()
    meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(meta.add_run("Student: [STUDENT NAME]\nStudent ID: [STUDENT ID]\nModule: CIS6005 Computational Intelligence\nAssessment: WRIT1\nAcademic year: 2025-2026"), size=11, color=MUTED)
    doc.add_page_break()

    doc.add_heading("Executive Summary", level=1)
    para(doc, f"This project builds a software-based virtual sensor that reconstructs atmospheric pressure from Mars Environmental Dynamics Analyzer (MEDA) observations in the Kaggle MEDA Virtual Sensor Recovery competition. The technical contribution is not leaderboard optimisation: it is a reproducible comparison of computational-intelligence techniques under severe temporal shift and an 8 GB CPU-only constraint. The original workflow trained seven model types and produced three accepted submissions, but its sampler read only the beginning of a single Parquet row group. Its apparent future MSE of 9.016 therefore conflicted with a public MSE of 5458.353. The rebuilt workflow streams all {len(raw):,} sampled observations across sols {int(raw.sol.min())}-{int(raw.sol.max())}, applies rolling-origin tuning, protects a final temporal test and packages {metadata['selected_model']} behind a Streamlit interface. On the untouched local test, the exported model obtained RMSE {final['rmse']:.3f} Pa, MAE {final['mae']:.3f} Pa and R2 {final['r2']:.3f}. The public-score disagreement remains a declared limitation and evidence that temporally convenient validation can be misleading.")

    doc.add_heading("1. Computational Intelligence and Traditional AI", level=1)
    para(doc, "Computational Intelligence (CI) describes adaptive, approximate problem-solving methods that learn useful behaviour from data or interaction. Neural networks, evolutionary computation, fuzzy systems and population-based optimisation are typical CI families. Their common strength is tolerance of noisy, nonlinear and partially observed environments. Traditional symbolic AI instead represents knowledge through explicit rules, logic and search. A rule-based Mars pressure system could encode statements about solar longitude or local time, but its performance would depend on complete expert rules and brittle thresholds. CI learns the mapping from heterogeneous rover measurements to pressure and can update when evidence changes.")
    para(doc, "The distinction is not absolute. This project combines empirical learning with physical representation: known periodic quantities are encoded as Fourier terms, while Ridge and ensemble learners estimate relationships from observations. This hybrid position is preferable to treating the system as either pure physics or a black box. Hao et al. (2022) argue that physics-informed machine learning can improve plausibility and data efficiency by embedding prior structure. Here, the diurnal, semidiurnal and seasonal encodings restrict representation rather than forcing a full atmospheric equation. The cost is that encoded priors can be incomplete; dust events, rover motion and sensor scheduling can still create patterns the simplified harmonics cannot explain.")
    para(doc, "CI is therefore suitable because the target is continuous, predictors are nonlinear and missingness is structural. It is also risky: flexible learners interpolate well but may fail when the prediction period lies beyond the training regime. The project treats validation design, reproducible artefacts and explicit limitations as part of intelligence rather than as administrative extras.")

    doc.add_heading("2. Literature Review", level=1)
    doc.add_heading("2.1 MEDA pressure and Jezero meteorology", level=2)
    para(doc, "Rodriguez-Manfredi et al. (2021) describe MEDA as an integrated suite measuring pressure, radiation, temperature, humidity, wind and dust. This supports multivariate recovery because radiative and geometric variables are physically related to atmospheric state. Jaakonaho et al. (2023) report that the MEDA pressure sensor is a calibrated low-pressure Barocap system and achieved high operational quality. Harri et al. (2024) estimate overall pressure uncertainty near 0.3%, approximately 1.8-2.4 Pa across 600-800 Pa. That scale matters when interpreting model RMSE: a numerically small error is not automatically equivalent to instrument accuracy.")
    para(doc, "Newman et al. (2022) analyse MEDA's first 250 sols and show that daily mean pressure rose from about 735 Pa during sols 15-20 to 761 Pa around sols 99-110, then declined towards 650 Pa by sol 250. They also identify diurnal, semidiurnal and terdiurnal thermal tides. Sanchez-Lavega et al. (2023) similarly connect pressure variability with tides, baroclinic waves, convection and a regional dust storm. These studies explain why raw mission sol and tree splits are inadequate for extrapolation: the seasonal curve changes direction and the daily signal is multi-harmonic. Consequently, the rebuilt model encodes solar longitude and three clock harmonics and validates by ordered sols.")
    doc.add_heading("2.2 Virtual sensing and alternative CI approaches", level=2)
    para(doc, "Masti, Bernardini and Bemporad (2021) define a lightweight ML virtual sensor that combines measurements with observer-derived features for parameter-varying systems. Their emphasis on bounded memory and online operation aligns with this project's deployment bundle. Perera et al. (2023) review AI-driven soft sensors and identify maintenance, drift, transparency and lifecycle monitoring as unresolved deployment problems. Those concerns are visible here: a saved model can run correctly while becoming scientifically unsuitable outside its seasonal range.")
    para(doc, "Hu et al. (2022) propose a spatio-temporal LSTM for sparse environmental sensor data and report performance comparable with Random Forest and XGBoost benchmarks. Vedavalli and Ch (2023) use a hierarchical LSTM with spatial-temporal correlation to recover erroneous IoT measurements. These deep approaches can represent sequences and shared sensor behaviour more explicitly than the row-wise MLP used here. However, they assume meaningful windows or neighbouring nodes. The competition table is handled as independent timestamped observations, and the available laptop constrains sequence construction and tuning. A shallow harmonic model is thus a defensible baseline, while XGBoost and CatBoost test nonlinear interactions.")
    para(doc, "Luo et al. (2023) survey adaptive, multiscale and deep soft sensors and show that no single architecture dominates all operating regimes. Hao et al. (2022) emphasise physical priors; the MEDA studies supply exactly such priors through seasonal carbon-dioxide cycling and thermal tides. The literature collectively favours a comparative design: regularised linear harmonics for transparent extrapolative structure, tree ensembles for nonlinear tabular effects, an MLP to test deep learning, and a residual hybrid to examine whether constrained nonlinear correction adds value.")

    doc.add_heading("3. Exploratory Data Analysis and Design Influence", level=1)
    para(doc, f"The raw training Parquet contains more than eight million records and one row group. A deterministic priority sampler streams the whole file and retains up to {metadata['rows_per_sol']:,} rows per sol. The resulting sample contains {len(raw):,} rows across {raw.sol.nunique()} unique sols. This corrects the original sampler, which selected from only the beginning and represented roughly sols 1-32.")
    add_figure(doc, FIGURES / "eda_pressure_distribution.png", "Figure 1. Pressure distribution in the complete sol-stratified sample.")
    para(doc, f"Figure 1 finding and implication. Pressure has mean {desc['mean']:.2f} Pa, median {desc['50%']:.2f} Pa, standard deviation {desc['std']:.2f} Pa and range {desc['min']:.2f}-{desc['max']:.2f} Pa. The distribution mixes seasonal and tidal states; therefore MAE/RMSE in pascals accompany MSE. Limitation: the histogram discards temporal order and cannot establish future performance.")
    add_figure(doc, FIGURES / "eda_missingness.png", "Figure 2. Highest missingness rates across MEDA fields.")
    para(doc, f"Figure 2 finding and implication. {missing.index[0]} is {missing.iloc[0]:.1f}% missing, and several channels are completely absent during the training period. Entirely missing variables contain no fitted information; partially observed fields require training-only imputation or missing-aware trees, plus an explicit missing-count feature. Limitation: percentages cannot distinguish planned instrument scheduling from failure, so missingness should not be interpreted causally.")
    add_figure(doc, FIGURES / "eda_pressure_by_sol.png", "Figure 3. Seasonal pressure movement by mission sol; shading shows within-sol standard deviation.")
    by_sol = raw.groupby("sol").PRESSURE.mean()
    para(doc, f"Figure 3 finding and implication. Mean pressure changes from {by_sol.iloc[0]:.2f} Pa at sol {int(by_sol.index[0])} to {by_sol.iloc[-1]:.2f} Pa at sol {int(by_sol.index[-1])}, with a curved rather than constant trend. Ordered splitting and cyclic solar-longitude features follow directly. Limitation: averaging hides vortices and within-sol pressure tides.")
    add_figure(doc, FIGURES / "eda_pressure_by_lmst.png", "Figure 4. Mean pressure by local mean solar hour.")
    para(doc, "Figure 4 finding and implication. Mean pressure changes systematically through the Martian day, supporting the thermal-tide evidence in Newman et al. (2022). LMST and LTST are therefore represented by first-, second- and third-order sine/cosine pairs. Limitation: aggregation across sols assumes a stable daily shape and conceals seasonal changes in tidal amplitude.")
    add_figure(doc, FIGURES / "eda_correlation.png", "Figure 5. Correlation structure for pressure and the strongest numerical associates.")
    para(doc, "Figure 5 finding and implication. Solar geometry, time and radiative measurements show the strongest marginal relationships with pressure, supporting a combined physical and environmental representation. Correlation is not used as automatic feature selection because shared time trends can create misleading association. Limitation: Pearson correlation measures linear dependence and cannot establish causality or nonlinear interaction.")

    doc.add_heading("4. System Architecture and ML Techniques", level=1)
    add_figure(doc, FIGURES / "architecture_training.png", "Figure 6. High-level training, evaluation and artefact architecture.")
    para(doc, "Figure 6 outcome. One shared feature builder feeds rolling-origin tuning, temporal validation, an untouched final test and a versioned bundle. This prevents the app from silently using different feature order or clock parsing. The trade-off is slower streaming and bounded training on the 8 GB laptop; representation improves while not every raw row is retained.")
    add_figure(doc, FIGURES / "architecture_inference.png", "Figure 7. High-level Streamlit inference architecture.")
    para(doc, "Figure 7 outcome. Both a guided form and CSV/Parquet upload pass through schema validation, the same transformation code and the selected model. The app returns pressure, warnings, previews and downloadable predictions. Unlike notebook-only competition examples, it exposes metadata and appropriate-use limitations. The trade-off is stricter input requirements; malformed clocks or absent core geometry are rejected rather than guessed.")
    para(doc, "Ridge supplies a transparent regularised harmonic baseline; Random Forest averages decorrelated trees; XGBoost and CatBoost fit sequential boosted trees; the MLP tests nonlinear neural representation; and the physics-residual model adds a conservative boosted correction to harmonic Ridge. The final application loads the validation-selected bundle, not a hard-coded algorithm name.")

    doc.add_heading("5. Implementation, Evaluation and Demonstration", level=1)
    para(doc, "Preprocessing converts numeric values to float32, removes PRESSURE and row_id, parses LMST/LTST, derives three clock harmonics, encodes solar longitude and rover angles cyclically, and counts missing fields. The first 70% of ordered unique sols form training, the next 15% validation and the latest 15% an untouched test. Three rolling-origin folds inside the first 85% tune Ridge alpha and XGBoost complexity. The final test is not referenced by tuning or selection.")
    para(doc, f"Seven candidates are evaluated with MSE, RMSE, MAE, R2, fit time and inference time. The lowest validation MSE was {validation.mse:.3f} from {validation.model}. After refitting the selected configuration on training plus validation, the protected test across sols {final['test_sol_min']}-{final['test_sol_max']} was accessed exactly once and produced MSE {final['mse']:.3f}, RMSE {final['rmse']:.3f} Pa, MAE {final['mae']:.3f} Pa and R2 {final['r2']:.3f}. Table 1 reports the validation comparison used for model selection; the protected-test result is reported only for the selected refitted model.")
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(c.add_run("Table 1. Multi-model temporal-validation comparison used for selection."), size=9, color=MUTED, italic=True)
    add_metric_table(doc, comparison)
    add_figure(doc, FIGURES / "evaluation_residuals.png", "Figure 8. Residuals on the untouched temporal test.")
    para(doc, "Figure 8 finding and implication. Residual dispersion and any visible trend reveal whether errors remain linked to pressure regime. This is more informative than a single average score and motivates error-by-sol monitoring. Limitation: held-out training-era sols are still nearer than Kaggle test sols, so the plot cannot certify remote seasonal extrapolation.")
    add_figure(doc, FIGURES / "evaluation_error_by_sol.png", "Figure 9. Mean squared error by held-out mission sol.")
    para(doc, "Figure 9 finding and implication. Error varies by sol rather than remaining stationary, reinforcing the need for temporal monitoring and periodic retraining. Limitation: sol-level averaging hides time-of-day error and may be unstable when observations per sol differ.")
    para(doc, "The Streamlit demonstration has three areas. Single prediction collects mission time, solar geometry, rover state and optional sensors before reporting pressure in pascals. Batch prediction validates CSV/Parquet data, predicts in 100,000-row chunks and returns row_id,PRESSURE. Model Information exposes the selected technique, protected-test metrics, coverage and warnings. Unit tests cover clock parsing, feature order, split disjointness, finite bundle prediction and missing required columns.")
    para(doc, "Kaggle evidence shows three accepted files: median baseline (7049.045), XGBoost (5470.251) and the original selected XGBoost (5458.353). Their SHA-256 hashes are retained. Leaderboard rank is irrelevant, but the discrepancy from local performance is a critical result. The supplied screenshot shows that no final/private candidate has yet been selected; this must be completed by the student before closure.")

    doc.add_heading("6. Critical Evaluation and Deep-Learning Suitability", level=1)
    para(doc, "The main strength is methodological correction. Full-file sol coverage, temporal folds, one preprocessing contract, saved trial histories and explicit bundle metadata make the work defensible. Multiple techniques expose trade-offs: ensembles model nonlinear interactions and missingness, harmonic Ridge encodes extrapolative cycles, and the hybrid tests limited residual correction. The working app demonstrates software engineering beyond a notebook.")
    para(doc, "The largest limitation is domain shift. Training and protected testing remain within the competition's labelled period, whereas Kaggle evaluates a later regime. The old public MSE proves that excellent interpolation can coexist with failed extrapolation. Published MEDA pressure climatology suggests broader seasonal structure, but this project deliberately avoids importing external target labels. Other limitations include bounded per-sol sampling, non-random missing sensors, unequal training sizes, absence of calibrated prediction intervals and dependence on organiser preprocessing.")
    para(doc, "Deep learning is not automatically superior. The MLP must infer cyclic structure from engineered rows, requires imputation and scaling, and consumes more tuning compute. Hu et al. (2022) and Vedavalli and Ch (2023) show that LSTM-based recovery can exploit spatial-temporal context, but this would require carefully constructed sequences, masking and a larger compute budget. A future temporal convolutional network or physics-informed recurrent model could learn changing tide amplitude while constraining seasonal behaviour. It should be evaluated with rolling origins, uncertainty intervals and ablation studies, not selected because it is deep.")
    para(doc, "Ethically, a Mars virtual sensor has lower direct human risk than medical or financial prediction, but fabricated certainty would still be harmful to scientific interpretation. The app labels itself educational and not flight-qualified; out-of-range inputs receive warnings; failures are not silently imputed for required fields. Reproducibility is supported through seeds, hashes, version metadata and preserved legacy artefacts. Emerging improvements include uncertainty-aware boosting, conformal intervals under drift, physics-informed neural time-series models and monitoring triggered by feature-range or residual change.")

    doc.add_heading("7. Conclusion", level=1)
    para(doc, "The project demonstrates CI as an adaptive but fallible engineering process. A diverse model comparison, bounded rolling-origin tuning and a usable application satisfy the practical requirements, while the corrected sampler and public-score critique show why validation design matters more than a superficially strong local score. The selected bundle is reproducible and operational for demonstration, not flight-qualified. The immediate administrative action is to choose and evidence a final/private Kaggle submission; the technical next step is uncertainty-aware, physics-informed sequence modelling across a wider seasonal record.")

    doc.add_heading("References", level=1)
    references = [
        "Hao, Z. et al. (2022) 'Physics-Informed Machine Learning: A Survey on Problems, Methods and Applications'. arXiv:2211.08064.",
        "Harri, A.-M. et al. (2024) 'Perseverance MEDA Atmospheric Pressure Observations - Initial Results', Journal of Geophysical Research: Planets, 129, e2023JE007880. https://doi.org/10.1029/2023JE007880",
        "Hu, Y. et al. (2022) 'A spatio-temporal LSTM model to forecast across multiple temporal and spatial scales', Ecological Informatics, 69, 101687. https://doi.org/10.1016/j.ecoinf.2022.101687",
        "Jaakonaho, I. et al. (2023) 'Pressure sensor for the Mars 2020 Perseverance rover', Planetary and Space Science, 239, 105815. https://doi.org/10.1016/j.pss.2023.105815",
        "Kaggle (2026) 'Mars Environmental Dynamics Analyzer (MEDA) Virtual Sensor Recovery'. Available at: https://www.kaggle.com/competitions/mars-environmental-dynamics-analyzer-meda-virtual-sensor-recovery (Accessed: 29 July 2026).",
        "Luo, Y. et al. (2023) 'Data-driven soft sensors in blast furnace ironmaking: a survey', Frontiers of Information Technology & Electronic Engineering, 24, pp. 327-354. https://doi.org/10.1631/FITEE.2200366",
        "Masti, D., Bernardini, D. and Bemporad, A. (2021) 'A machine-learning approach to synthesize virtual sensors for parameter-varying systems', European Journal of Control, 61, pp. 40-49. https://doi.org/10.1016/j.ejcon.2021.06.005",
        "Newman, C.E. et al. (2022) 'The diverse meteorology of Jezero crater over the first 250 sols of Perseverance on Mars', Nature Geoscience, 15, pp. 1058-1064. https://doi.org/10.1038/s41561-022-01084-0",
        "Perera, Y.S. et al. (2023) 'The role of artificial intelligence-driven soft sensors in advanced sustainable process industries: A critical review', Engineering Applications of Artificial Intelligence, 121, 105988. https://doi.org/10.1016/j.engappai.2023.105988",
        "Rodriguez-Manfredi, J.A. et al. (2021) 'The Mars Environmental Dynamics Analyzer, MEDA. A Suite of Environmental Sensors for the Mars 2020 Mission', Space Science Reviews, 217, 48. https://doi.org/10.1007/s11214-021-00816-9",
        "Sanchez-Lavega, A. et al. (2023) 'Mars 2020 Perseverance Rover Studies of the Martian Atmosphere Over Jezero From Pressure Measurements', Journal of Geophysical Research: Planets, 128, e2022JE007480. https://doi.org/10.1029/2022JE007480",
        "Vedavalli, P. and Ch, D. (2023) 'A Deep Learning Based Data Recovery Approach for Missing and Erroneous Data of IoT Nodes', Sensors, 23(1), 170. https://doi.org/10.3390/s23010170",
    ]
    for item in references:
        p = doc.add_paragraph(style="Normal"); p.paragraph_format.left_indent = Inches(.3); p.paragraph_format.first_line_indent = Inches(-.3)
        set_font(p.add_run(item), size=9.5)

    doc.add_page_break(); doc.add_heading("Appendix A - Kaggle Submission Evidence", level=1)
    para(doc, "The screenshot below records three successful model submissions and their public scores. The Select checkboxes are empty; a replacement screenshot must be captured after the student chooses the final/private candidate.")
    add_figure(doc, OUTPUTS / "evidence" / "kaggle_submissions.png", "Figure A1. Kaggle submissions supplied by the student on 29 July 2026.", width=6.4)
    doc.add_heading("Appendix B - Practical Demonstration", level=1)
    para(doc, "Run `.\\.venv\\Scripts\\streamlit.exe run app.py`, open the local URL, predict one guided observation, then upload `data/example_meda_input.csv` and download the generated row_id,PRESSURE file. The application displays the bundle version, temporal metrics and appropriate-use warning.")

    count = report_word_count(doc)
    if count > 4000:
        raise ValueError(f"Report exceeds 4,000 counted words: {count}")
    path = REPORT / "CIS6005_WRIT1_MEDA_Report.docx"
    doc.save(path)
    (OUTPUTS / "report_word_count.json").write_text(json.dumps({"counted_words_before_references": count,
        "limit": 4000, "student_identity_placeholders": True}, indent=2), encoding="utf-8")
    print(path); print("Counted words before references:", count)


if __name__ == "__main__":
    main()
